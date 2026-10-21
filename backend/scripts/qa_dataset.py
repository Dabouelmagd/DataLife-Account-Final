"""
Fictional test data for the QA company — every name, national ID, tax number,
phone and e-mail here is invented (e-mails use example.com, which never
delivers). Columns match the import screens' templates (api/import_data.py
TEMPLATES), so the same rows serve both the Excel files handed to testers
and scripts/seed_test_company.py --import.

Opening date for every balance: 2026-01-01.
"""

OPENING_DATE = "2026-01-01"

# ── employees: one per income-tax bracket band, from minimum wage up ──────────
# monthly basic → annual: 84k (bracket 2-3), … 1.8M (top bracket 27.5%)
EMPLOYEES = [
    # name, code, national_id, position, department, phone, hire_date, basic_salary
    ("منى عادل حسن",        "QA-001", "29503150101234", "عامل نظافة",      "الخدمات",          "01000000101", "2023-02-01",   7000),
    ("كريم سامي فؤاد",       "QA-002", "29811020101245", "أمين مخزن",        "المخازن",          "01000000102", "2022-07-15",   9000),
    ("هالة محمود رشدي",      "QA-003", "29207252101362", "سكرتيرة",          "الإدارة",          "01000000103", "2021-03-10",  12000),
    ("عمرو جلال ناصر",       "QA-004", "29001080101471", "محاسب",            "المالية",          "01000000104", "2020-09-01",  18000),
    ("ياسمين طارق شوقي",     "QA-005", "29409122101582", "أخصائي موارد بشرية", "الموارد البشرية", "01000000105", "2022-01-05",  25000),
    ("أحمد فتحي عبد الله",   "QA-006", "28806300101693", "مهندس مبيعات",     "المبيعات",         "01000000106", "2019-11-20",  35000),
    ("نورهان صلاح بدر",      "QA-007", "29112012101704", "رئيس حسابات",      "المالية",          "01000000107", "2018-05-14",  50000),
    ("مصطفى إبراهيم زكي",    "QA-008", "28502170101815", "مدير مبيعات",      "المبيعات",         "01000000108", "2017-08-01",  75000),
    ("داليا شريف منصور",     "QA-009", "28310052101926", "المدير المالي",    "المالية",          "01000000109", "2016-01-10", 110000),
    ("حسام الدين وائل كمال", "QA-010", "28004110101037", "المدير العام",     "الإدارة",          "01000000110", "2015-04-01", 150000),
]

# ── customers / suppliers: fictional names, invented 9-digit tax numbers ─────
# NOTE: for ETA pre-production submissions use the test taxpayer numbers the
# authority issues with your pre-production credentials — these are only for
# the import and the ERP's own validation.
CUSTOMERS = [
    # name, tax_id, email, phone, address, balance
    ("شركة النيل للتجارة (اختبار)",    "100000017", "nile@example.com",   "0223000001", "القاهرة — مدينة نصر", 120000),
    ("مؤسسة الدلتا للتوريدات (اختبار)", "100000025", "delta@example.com",  "0403000002", "طنطا",                 45000),
    ("شركة الأهرام للمقاولات (اختبار)", "100000033", "ahram@example.com",  "0233000003", "الجيزة — الدقي",       250000),
    ("مكتب الريادة الهندسي (اختبار)",   "100000041", "riyada@example.com", "0333000004", "الإسكندرية",           0),
    ("عميل نقدي — أفراد (اختبار)",      "",          "",                   "01000000201", "",                     0),
]

SUPPLIERS = [
    ("شركة الصفوة للخامات (اختبار)",   "200000016", "safwa@example.com",  "0223000011", "العاشر من رمضان", 180000),
    ("مصنع الشروق للتغليف (اختبار)",   "200000024", "shorouk@example.com", "0223000012", "6 أكتوبر",        60000),
    ("شركة المدار للنقل (اختبار)",     "200000032", "madar@example.com",  "0223000013", "القاهرة",          0),
]

# ── inventory: opening stock is posted by the import itself (Dr stock / Cr 213),
#    so stock is NOT repeated in the trial balance below ────────────────────────
INVENTORY = [
    # code, name, category, unit, quantity, cost_price, sale_price, min_stock
    ("QA-P001", "لابتوب أعمال 14 بوصة",  "أجهزة",    "قطعة", 40,  18000, 23500, 5),
    ("QA-P002", "شاشة 24 بوصة",          "أجهزة",    "قطعة", 60,   4200,  5600, 10),
    ("QA-P003", "طابعة ليزر",            "أجهزة",    "قطعة", 25,   6500,  8400, 5),
    ("QA-P004", "ورق تصوير A4 (كرتونة)", "مستلزمات", "كرتونة", 300,  950,  1250, 50),
    ("QA-P005", "حبر طابعة أسود",        "مستلزمات", "علبة",  200,  1100,  1450, 30),
]

# ── opening trial balance (balanced). Customers (131), suppliers (251) and
#    stock come from their own imports, so they are left out here. ────────────
OPENING_TB = [
    # code, name, debit, credit
    ("161",   "النقدية بالصندوق",                  750000,       0),
    ("162",   "النقدية بالبنوك الجارية",          1850000,       0),
    ("112",   "مباني وإنشاءات",                   3000000,       0),
    ("113",   "سيارات ووسائل نقل",                 800000,       0),
    ("115",   "أثاث وتجهيزات مكتبية",              250000,       0),
    ("116",   "أجهزة كمبيوتر وبرمجيات",            300000,       0),
    ("137",   "ضريبة المدخلات (VAT قابلة للاسترداد)", 45000,     0),
    ("22201", "مجمع إهلاك مباني وإنشاءات",              0,  300000),
    ("22203", "مجمع إهلاك سيارات ووسائل نقل",           0,  200000),
    ("22204", "مجمع إهلاك أثاث وأجهزة",                0,  150000),
    ("241",   "قروض بنكية طويلة الأجل",                 0, 1500000),
    ("252",   "أوراق الدفع",                            0,  120000),
    ("253",   "مصروفات مستحقة",                        0,   80000),
    ("255",   "التأمينات الاجتماعية المستحقة",          0,   45000),
    ("266",   "ضريبة كسب العمل المستقطعة",              0,   30000),
    ("211",   "رأس المال المدفوع",                      0, 4000000),
    ("213",   "الأرباح المرحلة",                        0,  570000),
]

# ── rows that MUST be refused — to check the import's validation ────────────
EMPLOYEES_WITH_ERRORS = [
    ("موظف رقم قومي ناقص",   "QA-901", "2950315010123",  "محاسب", "المالية", "01000000901", "2024-01-01", 9000),   # 13 digits
    ("موظف راتب سالب",       "QA-902", "29503150109999", "محاسب", "المالية", "01000000902", "2024-01-01", -5000),  # negative
    ("موظف بلا تاريخ تعيين",  "QA-903", "29503150108888", "محاسب", "المالية", "01000000903", "",           9000),   # missing date
    ("موظف رقم مكرر 1",      "QA-904", "29503150107777", "محاسب", "المالية", "01000000904", "2024-01-01", 9000),
    ("موظف رقم مكرر 2",      "QA-905", "29503150107777", "محاسب", "المالية", "01000000905", "2024-01-01", 9000),   # duplicate NID
]
CUSTOMERS_WITH_ERRORS = [
    ("عميل رقم ضريبي 8 أرقام", "12345678",  "bad8@example.com", "0223000901", "القاهرة", 0),
    ("عميل رصيد بلا تاريخ",    "100000099", "nodate@example.com", "0223000902", "القاهرة", 5000),
]
OPENING_TB_UNBALANCED = [
    ("162", "النقدية بالبنوك الجارية", 100000, 0),
    ("211", "رأس المال المدفوع",            0, 90000),   # off by 10,000
]


def _employees_rows(src):
    return [{"الاسم": n, "كود الموظف": c, "الرقم القومي": nid, "الوظيفة": p, "القسم": d,
             "البريد الإلكتروني": f"{c.lower()}@example.com", "الهاتف": ph,
             "تاريخ التعيين": h, "الراتب الأساسي": s} for n, c, nid, p, d, ph, h, s in src]


def _party_rows(src, with_date=True):
    out = []
    for n, t, e, ph, a, b in src:
        r = {"الاسم": n, "الرقم الضريبي": t, "البريد الإلكتروني": e, "الهاتف": ph,
             "العنوان": a, "الرصيد الافتتاحي": b}
        r["تاريخ الرصيد"] = OPENING_DATE if (with_date and b) else ""
        out.append(r)
    return out


def _tb_rows(src):
    return [{"كود الحساب": c, "اسم الحساب": n, "مدين": d or "", "دائن": cr or "",
             "تاريخ الرصيد": OPENING_DATE} for c, n, d, cr in src]


def files():
    """{filename: (sheet title, rows)} — the import order is the dict order."""
    return {
        "1-employees.xlsx":        ("الموظفين", _employees_rows(EMPLOYEES)),
        "2-customers.xlsx":        ("العملاء", _party_rows(CUSTOMERS)),
        "3-suppliers.xlsx":        ("الموردين", _party_rows(SUPPLIERS)),
        "4-inventory.xlsx":        ("المخزون", [{"كود الصنف": c, "الاسم": n, "الفئة": cat, "الوحدة": u,
                                                 "الكمية": q, "سعر التكلفة": cp, "سعر البيع": sp,
                                                 "الحد الأدنى": m, "تاريخ الرصيد": OPENING_DATE}
                                                for c, n, cat, u, q, cp, sp, m in INVENTORY]),
        "5-opening-balances.xlsx": ("الأرصدة الافتتاحية", _tb_rows(OPENING_TB)),
        "errors-employees.xlsx":   ("موظفين بأخطاء", _employees_rows(EMPLOYEES_WITH_ERRORS)),
        "errors-customers.xlsx":   ("عملاء بأخطاء", _party_rows(CUSTOMERS_WITH_ERRORS, with_date=False)),
        "errors-opening-unbalanced.xlsx": ("ميزان غير متوازن", _tb_rows(OPENING_TB_UNBALANCED)),
    }


IMPORT_ENDPOINTS = {
    "1-employees.xlsx": "/api/import/employees",
    "2-customers.xlsx": "/api/import/customers",
    "3-suppliers.xlsx": "/api/import/suppliers",
    "4-inventory.xlsx": "/api/import/inventory",
    "5-opening-balances.xlsx": "/api/import/opening-balances",
}


def write_xlsx(path, title, rows):
    """Text columns stay text (leading zeros in phones / IDs are kept)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True
    headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center")
    text_cols = {"الرقم القومي", "الهاتف", "الرقم الضريبي", "كود الحساب", "كود الموظف", "كود الصنف",
                 "تاريخ التعيين", "تاريخ الرصيد"}
    for r in rows:
        ws.append([r[h] for h in headers])
    for i, h in enumerate(headers, start=1):
        col = ws.cell(row=1, column=i).column_letter
        width = max(len(str(h)), *(len(str(r[h])) for r in rows)) + 4 if rows else 12
        ws.column_dimensions[col].width = min(width, 40)
        if h in text_cols:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                for c in row:
                    c.number_format = "@"
                    c.value = "" if c.value is None else str(c.value)
    ws.freeze_panes = "A2"
    wb.save(path)


def check_balanced():
    dr = sum(d for _, _, d, _ in OPENING_TB)
    cr = sum(c for _, _, _, c in OPENING_TB)
    assert dr == cr, (dr, cr)
    return dr


if __name__ == "__main__":
    import os
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "qa-data"
    os.makedirs(out, exist_ok=True)
    print("trial balance:", check_balanced())
    for name, (title, rows) in files().items():
        write_xlsx(os.path.join(out, name), title, rows)
        print("wrote", name, len(rows), "rows")
